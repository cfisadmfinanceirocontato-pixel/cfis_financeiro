"""
DIÁRIAS - Geração de recibos (versão adaptada)
"""
import streamlit as st
import pandas as pd
from datetime import datetime
import tempfile
from docx import Document

from src.database import MongoDB, DiariasDB
from src.utils import formatar_moeda_br, formatar_data_completa, limpar_cpf

st.set_page_config(page_title="Recibos de Diárias", layout="wide")

st.title("📋 Recibos de Diárias")

# ============================================================================
# FUNÇÕES AUXILIARES
# ============================================================================
@st.cache_data(ttl=300)
def carregar_termos():
    """Carrega termos do MongoDB"""
    colecao = MongoDB.get_collection('colaboradores')
    pipeline = [{'$group': {'_id': '$TERMO DE COLABORAÇÃO'}}, {'$sort': {'_id': 1}}]
    return [r['_id'] for r in colecao.aggregate(pipeline) if r['_id']]

@st.cache_data(ttl=300)
def carregar_funcionarios(termo):
    """Carrega funcionários de um termo"""
    colecao = MongoDB.get_collection('colaboradores')
    pipeline = [
        {'$match': {'TERMO DE COLABORAÇÃO': termo}},
        {'$group': {'_id': '$FUNCIONÁRIOS'}},
        {'$sort': {'_id': 1}}
    ]
    return [r['_id'] for r in colecao.aggregate(pipeline) if r['_id']]

@st.cache_data(ttl=300)
def buscar_dados_funcionario(termo, funcionario):
    """Busca CPF e cargo do funcionário"""
    colecao = MongoDB.get_collection('colaboradores')
    resultado = colecao.find_one({
        'TERMO DE COLABORAÇÃO': termo,
        'FUNCIONÁRIOS': funcionario
    })
    if resultado:
        return limpar_cpf(resultado.get('CPF', '')), resultado.get('CARGO', '')
    return '', ''

# ============================================================================
# INTERFACE
# ============================================================================
termos = carregar_termos()

with st.sidebar:
    st.header("🔍 Filtros")
    termo_filtro = st.selectbox("Termo:", ['Todos'] + termos)
    
    st.metric("Total Recibos", MongoDB.get_collection('diarias').count_documents({}))

# Formulário
with st.form("novo_recibo"):
    st.subheader("📝 Novo Recibo")
    
    col1, col2 = st.columns(2)
    with col1:
        termo = st.selectbox("Termo de Colaboração:", options=[''] + termos)
    with col2:
        instrumento = st.text_input("Instrumento:")
    
    if termo:
        funcionarios = carregar_funcionarios(termo)
        funcionario = st.selectbox("Funcionário:", options=[''] + funcionarios)
        
        if funcionario:
            cpf, cargo = buscar_dados_funcionario(termo, funcionario)
            col3, col4 = st.columns(2)
            with col3:
                cpf_input = st.text_input("CPF:", value=cpf)
            with col4:
                cargo_input = st.text_input("Cargo:", value=cargo)
    
    st.subheader("💰 Valores")
    qtd = st.number_input("Quantidade de diárias:", min_value=0.0, step=0.5, format="%.1f")
    valor_unitario = 140.0
    valor_total = qtd * valor_unitario
    
    col5, col6 = st.columns(2)
    with col5:
        st.metric("Valor Unitário", formatar_moeda_br(valor_unitario))
    with col6:
        st.metric("Valor Total", formatar_moeda_br(valor_total))
    
    data_recibo = st.date_input("Data do Recibo:", value=datetime.now())
    
    objetivo = st.text_area("Objetivo:")
    periodo = st.text_input("Período:", placeholder="01/02 a 03/02")
    
    template = st.file_uploader("Template (MODELO.docx)", type=['docx'])
    
    if st.form_submit_button("💾 Gerar Recibo"):
        if not termo or not funcionario or not template:
            st.error("Preencha os campos obrigatórios!")
        else:
            # Salvar template temporário
            with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
                tmp.write(template.read())
                template_path = tmp.name
            
            try:
                doc = Document(template_path)
                
                # Substituir placeholders
                replacements = {
                    "(FUNCIONÁRIO)": funcionario,
                    "(CARGO)": cargo_input,
                    "(CPF)": cpf_input,
                    "(VALOR)": formatar_moeda_br(valor_total),
                    "(QTD)": str(qtd),
                    "(OBJETIVO)": objetivo,
                    "(PERÍODO)": periodo,
                    "(DATA RECIBO)": formatar_data_completa(data_recibo)
                }
                
                for paragraph in doc.paragraphs:
                    for old, new in replacements.items():
                        if old in paragraph.text:
                            paragraph.text = paragraph.text.replace(old, new)
                
                # Salvar
                output_path = tempfile.NamedTemporaryFile(suffix='.docx', delete=False).name
                doc.save(output_path)
                
                with open(output_path, 'rb') as f:
                    docx_bytes = f.read()
                
                # Salvar no MongoDB
                dados = {
                    'Termo de Colaboração': termo,
                    'Instrumento': instrumento,
                    'Funcionário': funcionario,
                    'CPF': cpf_input,
                    'Cargo': cargo_input,
                    'Quantidade': str(qtd),
                    'Valor': formatar_moeda_br(valor_total),
                    'Período': periodo,
                    'Data Recibo': data_recibo.strftime('%d/%m/%Y')
                }
                
                DiariasDB.salvar_recibo(dados, arquivo_docx=docx_bytes)
                
                st.success("✅ Recibo gerado e salvo!")
                
                # Download
                st.download_button(
                    "📥 Download DOCX",
                    docx_bytes,
                    file_name=f"recibo_{funcionario.split()[0]}.docx"
                )
                
            except Exception as e:
                st.error(f"Erro: {e}")

# Lista recibos recentes
st.subheader("📋 Recibos Recentes")
recibos = DiariasDB.listar_recibos(limite=20)

if recibos:
    df = pd.DataFrame(recibos)
    if '_id' in df.columns:
        df = df.drop(columns=['_id'])
    st.dataframe(df, use_container_width=True)
